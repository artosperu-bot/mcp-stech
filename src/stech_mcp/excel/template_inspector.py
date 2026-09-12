from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from stech_mcp.excel.template_models import TemplateField, TemplateSchema


_ATTRIBUTE_ID_RE = re.compile(r"#([A-Za-z0-9_-]+)\b")
_AUXILIARY_SHEETS = {"OPCIONES", "MARCAS", "CATEGORÍAS", "CATEGORIAS"}


def _clean_values(values: Iterable[object]) -> tuple[str, ...]:
    output: list[str] = []
    for raw in values:
        value = str(raw or "").strip()
        if value and value not in output:
            output.append(value)
    return tuple(output)


def _first_column_values(workbook, sheet_names: tuple[str, ...]) -> tuple[str, ...]:
    for name in sheet_names:
        if name in workbook.sheetnames:
            ws = workbook[name]
            return _clean_values(ws.cell(row=row, column=1).value for row in range(2, ws.max_row + 1))
    return ()


def _inline_validation_values(worksheet, column: int, data_row: int) -> tuple[str, ...]:
    validations = getattr(getattr(worksheet, "data_validations", None), "dataValidation", []) or []
    for validation in validations:
        if str(getattr(validation, "type", "") or "").strip().lower() != "list":
            continue
        ranges = getattr(getattr(validation, "sqref", None), "ranges", []) or []
        applies = any(
            cell_range.min_col <= column <= cell_range.max_col
            and cell_range.min_row <= data_row <= cell_range.max_row
            for cell_range in ranges
        )
        if not applies:
            continue
        formula = str(getattr(validation, "formula1", "") or "").strip()
        if len(formula) >= 2 and formula.startswith('"') and formula.endswith('"'):
            return _clean_values(value.strip() for value in formula[1:-1].split(","))
    return ()


class TemplateInspector:
    """Inspect marketplace Excel templates without mutating their content."""

    @staticmethod
    def _data_sheet(workbook) -> str:
        for name in workbook.sheetnames:
            if name.strip().upper() not in _AUXILIARY_SHEETS:
                return name
        return workbook.sheetnames[0]

    @staticmethod
    def _header_row(worksheet) -> int:
        # V1 keeps discovery deterministic: choose the first row containing at
        # least two text headers. Real templates with leading title rows are
        # therefore supported without hard-coding a marketplace/category.
        for row in range(1, min(worksheet.max_row, 50) + 1):
            values = [worksheet.cell(row=row, column=column).value for column in range(1, worksheet.max_column + 1)]
            nonempty = [value for value in values if str(value or "").strip()]
            if len(nonempty) >= 2:
                return row
        return 1

    def inspect(self, path: str | Path) -> TemplateSchema:
        source = Path(path)
        workbook = load_workbook(source, data_only=False, read_only=False)
        try:
            data_sheet = self._data_sheet(workbook)
            worksheet = workbook[data_sheet]
            header_row = self._header_row(worksheet)
            fields: list[TemplateField] = []

            for column in range(1, worksheet.max_column + 1):
                raw = worksheet.cell(row=header_row, column=column).value
                header = str(raw or "").strip()
                if not header:
                    continue
                match = _ATTRIBUTE_ID_RE.search(header)
                fields.append(
                    TemplateField(
                        header=header,
                        column_index=column,
                        column_letter=get_column_letter(column),
                        attribute_id=match.group(1) if match else None,
                        required="*" in header,
                        allowed_values=_inline_validation_values(
                            worksheet,
                            column,
                            header_row + 1,
                        ),
                    )
                )

            brands = _first_column_values(workbook, ("Marcas", "MARCAS"))
            categories = _first_column_values(
                workbook,
                ("Categorías", "Categorias", "CATEGORÍAS", "CATEGORIAS"),
            )
            formula_cells = tuple(
                f"{data_sheet}!{cell.coordinate}"
                for row in worksheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            )
            return TemplateSchema(
                source_filename=source.name,
                data_sheet=data_sheet,
                sheet_names=tuple(workbook.sheetnames),
                fields=tuple(fields),
                brands=brands,
                categories=categories,
                formula_cells=formula_cells,
            )
        finally:
            workbook.close()
