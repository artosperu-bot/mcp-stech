from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

from openpyxl import load_workbook

from stech_mcp.excel.template_models import TemplateField, TemplateSchema


@dataclass(frozen=True)
class WriteReport:
    output_path: str
    updated_cells: tuple[str, ...]
    skipped_formula_cells: tuple[str, ...]


class ExcelWriter:
    """Write only explicitly requested data cells while preserving workbook structure."""

    @staticmethod
    def _field_index(schema: TemplateSchema) -> dict[str, TemplateField]:
        index: dict[str, TemplateField] = {}
        for field in schema.fields:
            keys = {
                field.stable_key,
                field.header,
                field.header.strip().casefold(),
                field.column_letter,
                field.column_letter.upper(),
            }
            if field.attribute_id:
                keys.add(field.attribute_id)
                keys.add(f"#{field.attribute_id}")
            for key in keys:
                index[str(key)] = field
        return index

    @staticmethod
    def _is_formula(value: Any) -> bool:
        return isinstance(value, str) and value.startswith("=")

    @staticmethod
    def _validate_closed_list(field: TemplateField, value: Any) -> None:
        if value is None or not field.allowed_values:
            return
        if str(value) not in field.allowed_values:
            raise ValueError(
                f"value {value!r} is not one of the allowed values for {field.header!r}: "
                f"{field.allowed_values!r}"
            )

    def write(
        self,
        source_path: str | Path,
        output_path: str | Path,
        schema: TemplateSchema,
        row_updates: Mapping[int, Mapping[str, Any]],
    ) -> WriteReport:
        source = Path(source_path)
        output = Path(output_path)
        if source.resolve() == output.resolve():
            raise ValueError("output_path must differ from source_path")

        workbook = load_workbook(source, data_only=False, read_only=False)
        updated: list[str] = []
        skipped_formulas: list[str] = []
        try:
            if schema.data_sheet not in workbook.sheetnames:
                raise ValueError(f"data sheet {schema.data_sheet!r} not found in workbook")
            worksheet = workbook[schema.data_sheet]
            field_index = self._field_index(schema)

            for row_number in sorted(row_updates):
                if int(row_number) < 1:
                    raise ValueError("row number must be positive")
                updates = row_updates[row_number]
                for raw_key, value in updates.items():
                    key = str(raw_key)
                    field = (
                        field_index.get(key)
                        or field_index.get(key.upper())
                        or field_index.get(key.casefold())
                    )
                    if field is None:
                        raise KeyError(f"template field {raw_key!r} was not discovered")

                    self._validate_closed_list(field, value)
                    cell = worksheet.cell(row=int(row_number), column=field.column_index)
                    cell_ref = f"{schema.data_sheet}!{cell.coordinate}"
                    if self._is_formula(cell.value):
                        skipped_formulas.append(cell_ref)
                        continue
                    if cell.value == value:
                        continue
                    cell.value = value
                    updated.append(cell_ref)

            output.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(output)
        finally:
            workbook.close()

        return WriteReport(
            output_path=str(output),
            updated_cells=tuple(updated),
            skipped_formula_cells=tuple(skipped_formulas),
        )
