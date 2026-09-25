from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from stech_mcp import server_authoritative as server


IMAGE_HEADERS = [
    "Imagen principal #IM1",
    "Imagen2 #IM2",
    "Imagen3 #IM3",
    "Imagen4 #IM4",
    "Imagen5 #IM5",
    "Imagen6 #IM6",
    "Imagen7 #IM7",
    "Imagen8 #IM8",
]


def _normalize(value: Any) -> str:
    return str(value or "").strip()


def _header_map(sheet, row: int = 4) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        value = _normalize(sheet.cell(row=row, column=column).value)
        if value:
            result[value] = column
    return result


def _find_sku_column(headers: dict[str, int]) -> int:
    for name, column in headers.items():
        if name.casefold().startswith("sku del vendedor"):
            return column
    raise ValueError("No se encontró la columna 'SKU del vendedor' en la fila 4")


def _output_path(input_path: Path, requested: str | None) -> Path:
    if requested:
        return Path(requested).expanduser().resolve()
    return input_path.with_name(f"{input_path.stem}_IMAGENES_FALABELLA.xlsx")


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Completa las columnas de imágenes de una plantilla Falabella usando "
            "las imágenes locales exactas de PC020 y URLs firmadas Cloudflare."
        )
    )
    parser.add_argument("input", help="Ruta al XLSX de Falabella")
    parser.add_argument("--output", default="")
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--sheet", default="Subir plantilla")
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    if not input_path.is_file():
        raise FileNotFoundError(str(input_path))
    output_path = _output_path(input_path, args.output or None)

    workbook = load_workbook(input_path)
    if args.sheet not in workbook.sheetnames:
        raise ValueError(f"No existe la hoja {args.sheet!r}")
    sheet = workbook[args.sheet]

    headers = _header_map(sheet, row=4)
    sku_column = _find_sku_column(headers)
    missing_headers = [name for name in IMAGE_HEADERS if name not in headers]
    if missing_headers:
        raise ValueError(f"Faltan columnas de imágenes: {missing_headers}")
    image_columns = [headers[name] for name in IMAGE_HEADERS]

    rows: list[dict[str, Any]] = []
    changed_cells = 0

    for row_number in range(5, sheet.max_row + 1):
        partnumber = _normalize(sheet.cell(row=row_number, column=sku_column).value).upper()
        if not partnumber:
            continue

        try:
            prepared = server.falabella_images_prepare(partnumber=partnumber, max_images=8)
            images = sorted(
                list(prepared.get("images") or []),
                key=lambda item: int(item.get("position") or 0),
            )
            # Falabella image columns are filled contiguously. If the local
            # inventory has a numbering gap (for example _01, _02, _04), the
            # visual order is preserved without leaving Image3 blank.
            ordered_urls = [
                str(item["url"])
                for item in images
                if item.get("url")
            ][:8]
            written = 0
            preserved = 0
            for column, url in zip(image_columns, ordered_urls, strict=False):
                cell = sheet.cell(row=row_number, column=column)
                if _normalize(cell.value) and not args.overwrite:
                    preserved += 1
                    continue
                cell.value = url
                cell.number_format = "@"
                changed_cells += 1
                written += 1

            rows.append(
                {
                    "row": row_number,
                    "partnumber": partnumber,
                    "state": prepared.get("state"),
                    "prepared_images": len(images),
                    "written_urls": written,
                    "preserved_existing": preserved,
                    "warning_count": prepared.get("warning_count", 0),
                    "errors": prepared.get("errors") or [],
                }
            )
        except Exception as exc:
            rows.append(
                {
                    "row": row_number,
                    "partnumber": partnumber,
                    "state": "ERROR",
                    "prepared_images": 0,
                    "written_urls": 0,
                    "preserved_existing": 0,
                    "warning_count": 0,
                    "errors": [f"{type(exc).__name__}: {exc}"],
                }
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    summary = {
        "ok": True,
        "input": str(input_path),
        "output": str(output_path),
        "sheet": args.sheet,
        "sku_count": len(rows),
        "changed_cells": changed_cells,
        "overwrite": bool(args.overwrite),
        "rows": rows,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
