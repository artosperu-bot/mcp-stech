from pathlib import Path

from openpyxl import Workbook

from stech_mcp.excel.template_inspector import TemplateInspector


def _make_book(path: Path, headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(headers)
    ws.append([None] * len(headers))

    brands = wb.create_sheet("Marcas")
    brands.append(["Marca"])
    brands.append(["LENOVO"])
    brands.append(["JBL"])

    categories = wb.create_sheet("Categorías")
    categories.append(["Categoría"])
    categories.append(["Portátiles|notebooks"])
    wb.save(path)


def test_template_inspector_exists():
    assert TemplateInspector is not None


def test_inspector_detects_stable_attribute_ids_required_fields_and_auxiliary_lists(tmp_path):
    source = tmp_path / "falabella.xlsx"
    _make_book(
        source,
        [
            "Sku vendedor * #1",
            "Código de barras * #56",
            "TipoDePortatil * #396393",
            "MemoriaRam #1540",
        ],
    )

    schema = TemplateInspector().inspect(source)

    assert schema.source_filename == "falabella.xlsx"
    assert schema.data_sheet == "Productos"
    assert set(schema.sheet_names) == {"Productos", "Marcas", "Categorías"}
    by_id = {field.attribute_id: field for field in schema.fields}
    assert by_id["56"].header == "Código de barras * #56"
    assert by_id["56"].required is True
    assert by_id["56"].column_letter == "B"
    assert by_id["1540"].required is False
    assert schema.brands == ("LENOVO", "JBL")
    assert schema.categories == ("Portátiles|notebooks",)


def test_inspector_uses_attribute_id_instead_of_column_position(tmp_path):
    source = tmp_path / "moved.xlsx"
    _make_book(
        source,
        [
            "MemoriaRam #1540",
            "TipoDePortatil * #396393",
            "Código de barras * #56",
            "Sku vendedor * #1",
        ],
    )

    schema = TemplateInspector().inspect(source)
    by_id = {field.attribute_id: field for field in schema.fields}

    assert set(by_id) == {"1", "56", "1540", "396393"}
    assert by_id["56"].column_letter == "C"


def test_inspector_supports_generic_template_without_attribute_ids(tmp_path):
    source = tmp_path / "smartphone.xlsx"
    _make_book(source, ["Part Number *", "Marca *", "RAM (GB)", "Sistema Operativo"])

    schema = TemplateInspector().inspect(source)

    assert schema.fields[0].header == "Part Number *"
    assert schema.fields[0].required is True
    assert schema.fields[0].attribute_id is None
    assert schema.fields[2].header == "RAM (GB)"
