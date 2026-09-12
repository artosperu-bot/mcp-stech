from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from stech_mcp.excel.excel_writer import ExcelWriter
from stech_mcp.excel.template_inspector import TemplateInspector


def _make_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Productos"
    ws.append(["Part Number * #1", "Gama * #2", "Calculado #3"])
    ws["A1"].font = Font(bold=True)
    ws["C2"] = "=LEN(A2)"
    validation = DataValidation(type="list", formula1='"Baja,Media,Alta"')
    ws.add_data_validation(validation)
    validation.add("B2:B100")
    aux = wb.create_sheet("Opciones")
    aux.append(["No tocar"])
    aux.append(["valor auxiliar"])
    wb.save(path)


def test_writer_changes_only_requested_data_cells_and_preserves_structure(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _make_template(source)
    schema = TemplateInspector().inspect(source)

    report = ExcelWriter().write(
        source,
        output,
        schema,
        {2: {"1": "82YU00XYLM", "2": "Media", "3": 99}},
    )

    assert report.updated_cells == ("Productos!A2", "Productos!B2")
    assert report.skipped_formula_cells == ("Productos!C2",)

    wb = load_workbook(output, data_only=False)
    try:
        assert wb.sheetnames == ["Productos", "Opciones"]
        ws = wb["Productos"]
        assert [ws.cell(row=1, column=i).value for i in range(1, 4)] == [
            "Part Number * #1",
            "Gama * #2",
            "Calculado #3",
        ]
        assert ws["A2"].value == "82YU00XYLM"
        assert ws["B2"].value == "Media"
        assert ws["C2"].value == "=LEN(A2)"
        assert ws["A1"].font.bold is True
        assert len(ws.data_validations.dataValidation) == 1
        assert wb["Opciones"]["A2"].value == "valor auxiliar"
    finally:
        wb.close()


def test_writer_rejects_value_outside_closed_list(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _make_template(source)
    schema = TemplateInspector().inspect(source)

    try:
        ExcelWriter().write(source, output, schema, {2: {"2": "Premium"}})
    except ValueError as exc:
        assert "allowed values" in str(exc)
    else:
        raise AssertionError("closed-list violation must fail")


def test_writer_never_changes_source_file(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    _make_template(source)
    before = source.read_bytes()
    schema = TemplateInspector().inspect(source)

    ExcelWriter().write(source, output, schema, {2: {"1": "PN-1"}})

    assert source.read_bytes() == before
    assert output.exists()
